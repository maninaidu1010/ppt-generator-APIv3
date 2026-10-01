"""
PPT Generator API  (Flask + python-pptx)
========================================

Drop-in replacement for the original app.py. Same routes (/createppt, /download/<file>),
same minimum payload, but the decks are now designed instead of plain text:

  * dark title / section / closing slides and light content slides
  * auto agenda, numbered cards, icon-style rows, KPI stat cards, process flow,
    side-by-side comparison, native (editable) charts and tables
  * consistent theme (colours + fonts), slide numbers, speaker notes
  * automatic layout selection and text fitting, so even the old payload
    {"presentationTitle": "...", "slides": [{"title": "...", "content": ["...", ...]}]}
    produces a varied, professional deck

Request body for POST /createppt  (everything except slides is optional)
-----------------------------------------------------------------------
{
  "presentationTitle": "Service Desk Performance Review",
  "subtitle": "Q3 FY26 summary",
  "organization": "IT Service Management",
  "presenter": "Mani", "date": "October 2026",
  "theme": "midnight | forest | ocean | charcoal | berry",
  "footer": "text for the footer (default: presentation title, "" hides it)",
  "agenda": true, "closing": true,
  "closingTitle": "Thank You", "closingSubtitle": "Questions & discussion",
  "slides": [
    {
      "title": "Slide title",
      "layout": "auto | cards | rows | split | stats | process | compare | chart | table | section | spotlight",
      "content": ["Heading: supporting sentence", "Another point", ...],
      "notes": "Speaker notes",
      "subtitle": "used by section slides",
      "stats":  [{"value": "96%", "label": "SLA compliance"}],
      "steps":  [{"title": "Log", "description": "Ticket captured"}],
      "left":   {"heading": "Before", "points": ["..."]},
      "right":  {"heading": "After",  "points": ["..."]},
      "chart":  {"type": "column|bar|line|pie|doughnut|stacked|area",
                 "title": "optional", "categories": ["Jan","Feb"],
                 "series": [{"name": "Tickets", "values": [120, 98]}]},
      "table":  {"headers": ["Priority","Target"], "rows": [["P1","1 hr"]]}
    }
  ]
}
"""

import json
import math
import os
import re
import time
import uuid

from flask import Flask, jsonify, request, send_from_directory
from lxml import etree
from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.dml.color import RGBColor
from pptx.enum.chart import XL_CHART_TYPE, XL_LABEL_POSITION, XL_LEGEND_POSITION, XL_MARKER_STYLE
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, MSO_AUTO_SIZE, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Emu, Inches, Pt

app = Flask(__name__)

OUTPUT_FOLDER = "generated_ppts"
os.makedirs(OUTPUT_FOLDER, exist_ok=True)

MAX_CONTENT_SLIDES = int(os.environ.get("MAX_SLIDES", 12))
RETENTION_HOURS = int(os.environ.get("RETENTION_HOURS", 24))

# =========================
# THEMES
# =========================
# dark   : title/closing background and headings
# accent : numbers, icons, chart highlight
# card   : light tint behind content blocks

THEMES = {
    "midnight": dict(dark="14213D", accent="1B9AAA", card="F1F5F9", text="1F2937", muted="64748B",
                     chart=["1B9AAA", "14213D", "F2A541", "7C8DB5", "E4572E", "4CAF93"]),
    "forest":   dict(dark="1E3D2F", accent="3A9D5D", card="F1F6F2", text="1F2937", muted="5F6B63",
                     chart=["3A9D5D", "1E3D2F", "E0A526", "7FB88F", "C8553D", "5B7F95"]),
    "ocean":    dict(dark="0B3C5D", accent="328CC1", card="EFF5FA", text="1F2937", muted="5F7184",
                     chart=["328CC1", "0B3C5D", "D9B310", "6FB1D8", "D1495B", "3E8E7E"]),
    "charcoal": dict(dark="262B33", accent="F26B38", card="F3F4F6", text="1F2937", muted="6B7280",
                     chart=["F26B38", "262B33", "3D8BFD", "9AA5B1", "2FA37A", "E5B93C"]),
    "berry":    dict(dark="4A1942", accent="C44F7B", card="F8F1F5", text="1F2937", muted="6F5B69",
                     chart=["C44F7B", "4A1942", "E8A33D", "8E6C9B", "3E8E7E", "4F6D9A"]),
}
THEME_ALIASES = {"blue": "ocean", "green": "forest", "orange": "charcoal", "purple": "berry",
                 "teal": "midnight", "navy": "midnight", "default": "midnight"}

HEAD_FONT = "Cambria"   # safe fonts: ship with Office, render true-to-width
BODY_FONT = "Calibri"
ON_DARK = "D5DEEA"
ON_DARK_MUTED = "9FB0C7"

# content area (inches) on a 13.333 x 7.5 slide
AX, AY, AW, AH = 0.7, 1.8, 11.93, 4.9


# =========================
# SMALL HELPERS
# =========================

def rgb(hex_color):
    return RGBColor.from_string(hex_color)


def clean(text):
    """Normalise one piece of text coming from an LLM (markdown, bullets, spacing)."""
    if text is None:
        return ""
    t = str(text).replace("**", "").replace("__", "").replace("`", "")
    t = re.sub(r"^\s*(?:[-*•✓✔▪●]+|\d{1,2}[.)])\s+", "", t)
    return re.sub(r"[ \t]+", " ", t).strip()


def truncate_to(text, max_chars):
    if len(text) <= max_chars:
        return text
    cut = text[: max(1, max_chars - 1)].rsplit(" ", 1)[0]
    return cut.rstrip(",;:.-– ") + "…"


def split_item(text):
    """'Heading: description' -> ('Heading', 'description'). Short text becomes a heading."""
    m = re.match(r"^(.{2,60}?)\s*(?::|\s[-–—]\s)\s*(.+)$", text)
    if m:
        return m.group(1).strip(), m.group(2).strip()
    if len(text) <= 45:
        return text, ""
    return "", text


_NUM = r"[~≈<>$₹€£]?\s?\d[\d,\.]*\s?(?:%|[KkMmBb]\b|x\b|\+|[A-Za-z]{1,5}\b)?"
_STAT_A = re.compile(rf"^({_NUM})\s*[-–—:|]\s*(.+)$")
_STAT_B = re.compile(rf"^(.+?)\s*[:–—-]\s*({_NUM})\s*$")
_YEAR = re.compile(r"(19|20)\d\d")


def parse_stat(text):
    m = _STAT_A.match(text)
    if m and len(m.group(1).strip()) <= 12 and not _YEAR.fullmatch(m.group(1).strip()):
        return m.group(1).strip(), m.group(2).strip()
    m = _STAT_B.match(text)
    if m and len(m.group(2).strip()) <= 12 and not _YEAR.fullmatch(m.group(2).strip()):
        return m.group(2).strip(), m.group(1).strip()
    return None


def normalize_items(content, limit=8):
    if content is None:
        return []
    if isinstance(content, str):
        content = [c for c in re.split(r"[\r\n]+", content) if c.strip()]
    if not isinstance(content, (list, tuple)):
        content = [content]
    items = []
    for c in content:
        if isinstance(c, dict):
            head = clean(c.get("heading") or c.get("title") or "")
            body = clean(c.get("text") or c.get("description") or c.get("body") or "")
            c = f"{head}: {body}" if head and body else (head or body)
        c = clean(c)
        if c:
            items.append(c)
    return items[:limit]


# ---- text fitting (python-pptx cannot measure text, so we estimate it) ----

def _line_count(pars, pt, width_in, cw):
    cpl = max(1, int(width_in * 72 / (pt * cw)))
    return sum(max(1, math.ceil(len(p) / cpl)) for p in pars)


def fit_size(pars, w, h, max_pt, min_pt, bold=False, gap=0.0, cw=None):
    cw = cw or (0.54 if bold else 0.50)
    pt = int(max_pt)
    while pt >= min_pt:
        need = _line_count(pars, pt, w, cw) * pt * 1.2 / 72 + gap * pt / 72 * (len(pars) - 1)
        if need <= h:
            return pt, True
        pt -= 1
    return int(min_pt), False


# =========================
# LOW-LEVEL SHAPE / TEXT BUILDERS
# =========================

def fill_tf(tf, pars, size, color, font, bold=False, italic=False, align="l", anchor="t",
            bullets=False, gap=0.0, bullet_color=None):
    tf.word_wrap = True
    tf.auto_size = MSO_AUTO_SIZE.NONE
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = {"t": MSO_ANCHOR.TOP, "m": MSO_ANCHOR.MIDDLE, "b": MSO_ANCHOR.BOTTOM}[anchor]
    for i, par in enumerate(pars):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = {"l": PP_ALIGN.LEFT, "c": PP_ALIGN.CENTER, "r": PP_ALIGN.RIGHT}[align]
        if gap and i < len(pars) - 1:
            p.space_after = Pt(size * gap)
        run = p.add_run()
        run.text = par
        f = run.font
        f.name, f.size, f.bold, f.italic = font, Pt(size), bold, italic
        f.color.rgb = rgb(color)
        if bullets:
            pPr = p._p.get_or_add_pPr()
            pPr.set("marL", str(int(Inches(0.28))))
            pPr.set("indent", str(-int(Inches(0.28))))
            if bullet_color:
                bc = etree.SubElement(pPr, qn("a:buClr"))
                etree.SubElement(bc, qn("a:srgbClr")).set("val", bullet_color)
            etree.SubElement(pPr, qn("a:buChar")).set("char", "•")


def no_shadow(shape):
    try:
        shape.shadow.inherit = False
        ref = shape._element.find(qn("p:style"))
        if ref is not None:
            ref.find(qn("a:effectRef")).set("idx", "0")
    except Exception:
        pass


def add_box(slide, x, y, w, h, fill=None, shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.14, alpha=None):
    shp = slide.shapes.add_shape(shape, Inches(x), Inches(y), Inches(w), Inches(h))
    no_shadow(shp)
    if fill:
        shp.fill.solid()
        shp.fill.fore_color.rgb = rgb(fill)
        if alpha is not None:  # alpha = opacity in percent
            clr = shp._element.spPr.find(qn("a:solidFill")).find(qn("a:srgbClr"))
            etree.SubElement(clr, qn("a:alpha")).set("val", str(int(alpha * 1000)))
    else:
        shp.fill.background()
    shp.line.fill.background()
    if shape == MSO_SHAPE.ROUNDED_RECTANGLE:
        shp.adjustments[0] = min(0.5, radius / min(w, h))
    return shp


# =========================
# DECK BUILDER
# =========================

class Deck:
    def __init__(self, title, theme="midnight", footer=None):
        self.prs = Presentation()
        self.prs.slide_width, self.prs.slide_height = Emu(12192000), Emu(6858000)  # 16:9
        key = str(theme or "midnight").lower()
        self.T = THEMES.get(THEME_ALIASES.get(key, key), THEMES["midnight"])
        self.title = title
        self.footer_text = title if footer is None else footer
        self.page = 0
        self.variety = 0
        self.section_no = 0

    # ---------- primitives ----------

    def text(self, slide, x, y, w, h, text, size=16, bold=False, color=None, font=BODY_FONT, align="l",
             anchor="t", italic=False, bullets=False, min_size=None, gap=0.0, cw=None, bullet_color=None):
        pars = [text] if isinstance(text, str) else [t for t in text]
        pars = [p for p in pars if p]
        if not pars:
            return None
        inner_w = w - (0.28 if bullets else 0)
        min_size = min_size or max(10, int(size * 0.7))
        sz, ok = fit_size(pars, inner_w, h, size, min_size, bold, gap, cw)
        tries = 0
        while not ok and tries < 6:  # last resort: shorten text instead of overflowing
            pars = [truncate_to(p, int(len(p) * 0.85)) for p in pars]
            sz, ok = fit_size(pars, inner_w, h, size, min_size, bold, gap, cw)
            tries += 1
        tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
        fill_tf(tb.text_frame, pars, sz, color or self.T["text"], font, bold, italic, align, anchor,
                bullets, gap, bullet_color)
        return tb

    def circle(self, slide, x, y, d, label, fill=None, color="FFFFFF", size=18):
        c = add_box(slide, x, y, d, d, fill=fill or self.T["accent"], shape=MSO_SHAPE.OVAL)
        fill_tf(c.text_frame, [str(label)], size, color, BODY_FONT, bold=True, align="c", anchor="m")
        return c

    def set_title(self, slide, text, x, y, w, h, size, color, anchor="m", min_size=24):
        ph = slide.shapes.title
        ph.left, ph.top, ph.width, ph.height = Inches(x), Inches(y), Inches(w), Inches(h)
        sz, ok = fit_size([text], w, h, size, min_size, True, cw=0.58)
        if not ok:
            text = truncate_to(text, 80)
        fill_tf(ph.text_frame, [text], sz, color, HEAD_FONT, bold=True, anchor=anchor)

    def new_slide(self, dark=False):
        slide = self.prs.slides.add_slide(self.prs.slide_layouts[5])  # "Title Only" keeps a real title
        self.page += 1
        bg = slide.background.fill
        bg.solid()
        bg.fore_color.rgb = rgb(self.T["dark"] if dark else "FFFFFF")
        return slide

    def content_slide(self, title):
        slide = self.new_slide()
        self.set_title(slide, title, AX, 0.5, AW, 1.0, 32, self.T["dark"])
        if self.footer_text:
            self.text(slide, AX, 6.95, 9.0, 0.3, self.footer_text, size=10, color=self.T["muted"], min_size=9)
        self.text(slide, 11.63, 6.95, 1.0, 0.3, str(self.page), size=10, color=self.T["muted"], align="r")
        return slide

    def notes(self, slide, text):
        if text:
            slide.notes_slide.notes_text_frame.text = str(text)

    def decor(self, slide, small=False):
        """Circle motif used on dark slides (also echoed by numbered circles on content slides)."""
        a = self.T["accent"]
        if small:
            add_box(slide, 9.4, 1.6, 3.4, 3.4, fill=a, shape=MSO_SHAPE.OVAL, alpha=22)
            add_box(slide, 10.5, 2.7, 2.0, 2.0, fill="FFFFFF", shape=MSO_SHAPE.OVAL, alpha=10)
            add_box(slide, 9.1, 4.5, 0.55, 0.55, fill=a, shape=MSO_SHAPE.OVAL)
        else:
            add_box(slide, 7.9, 1.0, 5.0, 5.0, fill=a, shape=MSO_SHAPE.OVAL, alpha=22)
            add_box(slide, 9.5, 2.6, 2.9, 2.9, fill="FFFFFF", shape=MSO_SHAPE.OVAL, alpha=10)
            add_box(slide, 7.6, 5.2, 0.8, 0.8, fill=a, shape=MSO_SHAPE.OVAL)

    # ---------- dark slides ----------

    def title_slide(self, subtitle="", eyebrow="", meta=""):
        slide = self.new_slide(dark=True)
        self.decor(slide)
        if eyebrow:
            self.text(slide, 0.9, 1.9, 7.0, 0.4, eyebrow.upper(), size=14, bold=True, color=self.T["accent"])
        self.set_title(slide, self.title, 0.9, 2.4, 7.0, 2.3, 44, "FFFFFF", anchor="t", min_size=28)
        if subtitle:
            self.text(slide, 0.9, 4.9, 6.8, 0.9, subtitle, size=20, color=ON_DARK, min_size=14)
        if meta:
            self.text(slide, 0.9, 6.2, 6.8, 0.4, meta, size=14, color=ON_DARK_MUTED, min_size=11)
        return slide

    def section_slide(self, title, subtitle=""):
        self.section_no += 1
        slide = self.new_slide(dark=True)
        self.decor(slide, small=True)
        self.text(slide, 0.9, 1.9, 4.0, 1.4, f"{self.section_no:02d}", size=72, bold=True,
                  color=self.T["accent"], font=HEAD_FONT, min_size=48)
        self.set_title(slide, title, 0.9, 3.4, 8.0, 1.4, 40, "FFFFFF", anchor="t", min_size=26)
        if subtitle:
            self.text(slide, 0.9, 5.0, 7.6, 1.0, subtitle, size=18, color=ON_DARK, min_size=13)
        return slide

    def closing_slide(self, title="Thank You", subtitle=""):
        slide = self.new_slide(dark=True)
        self.decor(slide)
        self.set_title(slide, title, 0.9, 2.6, 7.0, 1.4, 54, "FFFFFF", anchor="t", min_size=32)
        if subtitle:
            self.text(slide, 0.9, 4.2, 6.8, 1.0, subtitle, size=20, color=ON_DARK, min_size=14)
        return slide

    # ---------- light content layouts ----------

    def rows(self, slide, items, x, y, w, h, cols=1, max_row_h=1.3, start=1, gap=0.2):
        """Numbered circle + heading + description inside light cards."""
        n = len(items)
        per = math.ceil(n / cols)
        col_w = (w - gap * (cols - 1)) / cols
        row_h = min(max_row_h, (h - gap * (per - 1)) / per)
        for i, (head, desc) in enumerate(items):
            c, r = divmod(i, per)
            rx, ry = x + c * (col_w + gap), y + r * (row_h + gap)
            add_box(slide, rx, ry, col_w, row_h, fill=self.T["card"])
            d = min(0.6, row_h - 0.3)
            self.circle(slide, rx + 0.25, ry + (row_h - d) / 2, d, start + i, size=16)
            tx = rx + 0.25 + d + 0.25
            tw = col_w - (tx - rx) - 0.3
            if head and desc:
                th = row_h - 0.24
                hh = max(0.35, th * 0.4)
                self.text(slide, tx, ry + 0.12, tw, hh, head, size=18, bold=True, color=self.T["dark"],
                          anchor="b", min_size=13)
                self.text(slide, tx, ry + 0.12 + hh + 0.03, tw, th - hh - 0.03, desc, size=14, min_size=11)
            else:
                self.text(slide, tx, ry + 0.1, tw, row_h - 0.2, head or desc, size=18 if head else 16,
                          bold=bool(head), color=self.T["dark"] if head else None, anchor="m", min_size=12)

    def cards(self, slide, items):
        n = len(items)
        cols = n if n <= 3 else (2 if n == 4 else 3)
        rows = math.ceil(n / cols)
        gap = 0.3
        cw = (AW - gap * (cols - 1)) / cols
        ch = min(3.4, (AH - gap * (rows - 1)) / rows)
        y0 = AY + (AH - (rows * ch + gap * (rows - 1))) / 2 if rows == 1 else AY
        for i, (head, desc) in enumerate(items):
            r, c = divmod(i, cols)
            cx, cy = AX + c * (cw + gap), y0 + r * (ch + gap)
            add_box(slide, cx, cy, cw, ch, fill=self.T["card"])
            self.circle(slide, cx + 0.3, cy + 0.3, 0.6, i + 1, size=18)
            if rows == 1:   # tall card: heading under the circle
                hy, hh = cy + 1.15, min(0.9, ch * 0.22)
                tx, tw = cx + 0.3, cw - 0.6
            else:           # compact card: heading beside the circle
                hy, hh = cy + 0.3, 0.6
                tx, tw = cx + 1.1, cw - 1.4
            if head and desc:
                self.text(slide, tx, hy, tw, hh, head, size=21 if rows == 1 else 19, bold=True,
                          color=self.T["dark"], anchor="m" if rows > 1 else "t", min_size=14)
                dy = hy + hh + 0.1 if rows == 1 else cy + 1.1
                self.text(slide, cx + 0.3, dy, cw - 0.6, cy + ch - dy - 0.3, desc,
                          size=16 if rows == 1 else 14, min_size=11)
            else:
                dy = cy + 1.15 if rows == 1 else cy + 1.1
                self.text(slide, cx + 0.3, dy, cw - 0.6, cy + ch - dy - 0.3, head or desc, size=18,
                          bold=bool(head), color=self.T["dark"] if head else None, min_size=12)

    def split(self, slide, items):
        head, desc = split_item(items[0])
        add_box(slide, AX, AY, 4.3, AH, fill=self.T["dark"])
        self.text(slide, AX + 0.4, AY + 0.4, 3.5, 0.3, "KEY POINT", size=12, bold=True, color=self.T["accent"])
        self.text(slide, AX + 0.4, AY + 0.85, 3.5, 2.3, head or desc, size=26, bold=True, color="FFFFFF",
                  font=HEAD_FONT, min_size=16, cw=0.58)
        if head and desc:
            self.text(slide, AX + 0.4, AY + 3.3, 3.5, 1.3, desc, size=15, color=ON_DARK, min_size=11)
        rest = [split_item(i) for i in items[1:6]]
        self.rows(slide, rest, AX + 4.7, AY, AW - 4.7, AH, cols=1, max_row_h=1.3, start=2)

    def spotlight(self, slide, text):
        add_box(slide, AX, 2.0, AW, 3.8, fill=self.T["dark"])
        add_box(slide, AX + 0.5, 2.5, 0.5, 0.5, fill=self.T["accent"], shape=MSO_SHAPE.OVAL)
        self.text(slide, AX + 0.5, 3.3, AW - 1.0, 2.2, text, size=30, color="FFFFFF", font=HEAD_FONT,
                  min_size=18, cw=0.55)

    def stats(self, slide, stats, extra_items=None):
        n = len(stats)
        gap = 0.3
        cw = (AW - gap * (n - 1)) / n
        ch = 2.4 if extra_items else 2.6
        y0 = AY if extra_items else AY + 0.9
        for i, (val, label) in enumerate(stats):
            cx = AX + i * (cw + gap)
            add_box(slide, cx, y0, cw, ch, fill=self.T["card"])
            self.text(slide, cx + 0.3, y0 + 0.3, cw - 0.6, 1.1, val, size=48, bold=True, color=self.T["accent"],
                      font=HEAD_FONT, anchor="m", min_size=24, cw=0.58)
            self.text(slide, cx + 0.3, y0 + 1.5, cw - 0.6, ch - 1.7, label, size=16, min_size=11)
        if extra_items:
            by = y0 + ch + 0.3
            add_box(slide, AX, by, AW, AY + AH - by, fill=self.T["card"])
            self.text(slide, AX + 0.4, by + 0.25, AW - 0.8, AY + AH - by - 0.5, extra_items[:3], size=16,
                      bullets=True, gap=0.4, bullet_color=self.T["accent"], anchor="m", min_size=12)

    def process(self, slide, steps):
        n = len(steps)
        if n > 5:
            return self.rows(slide, steps, AX, AY, AW, AH, cols=2, max_row_h=1.3)
        gap = 0.3
        cw = (AW - gap * (n - 1)) / n
        d = 0.8
        line = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(AX + cw / 2), Inches(AY + d / 2),
                                          Inches(AX + AW - cw / 2), Inches(AY + d / 2))
        line.line.color.rgb = rgb(self.T["accent"])
        line.line.width = Pt(2)
        top = AY + d / 2
        has_desc = any(desc for _, desc in steps)
        for i, (head, desc) in enumerate(steps):
            cx = AX + i * (cw + gap)
            card_h = AH - d / 2 if has_desc else 2.6
            add_box(slide, cx, top, cw, card_h, fill=self.T["card"])
            self.circle(slide, cx + cw / 2 - d / 2, AY, d, i + 1, size=22)
            if has_desc:
                self.text(slide, cx + 0.25, top + 0.65, cw - 0.5, 0.9, head or desc, size=19, bold=True,
                          color=self.T["dark"], align="c", anchor="m", min_size=13)
                if head and desc:
                    self.text(slide, cx + 0.25, top + 1.7, cw - 0.5, card_h - 1.95, desc, size=15, align="c",
                              min_size=11)
            else:
                self.text(slide, cx + 0.25, top + 0.55, cw - 0.5, card_h - 0.8, head or desc, size=18, bold=True,
                          color=self.T["dark"], align="c", anchor="m", min_size=12)

    def compare(self, slide, left, right):
        gap = 0.4
        cw = (AW - gap) / 2
        body_h = min(AH - 1.0, 0.8 * max(len(left["points"]), len(right["points"]), 1) + 1.0)
        for k, (side, fill) in enumerate(((left, self.T["dark"]), (right, self.T["accent"]))):
            cx = AX + k * (cw + gap)
            add_box(slide, cx, AY, cw, 0.8, fill=fill)
            self.text(slide, cx + 0.4, AY + 0.05, cw - 0.8, 0.7, side["heading"], size=22, bold=True,
                      color="FFFFFF", anchor="m", min_size=14)
            add_box(slide, cx, AY + 1.0, cw, body_h, fill=self.T["card"])
            self.text(slide, cx + 0.4, AY + 1.3, cw - 0.8, body_h - 0.6, side["points"][:6], size=18, bullets=True,
                      gap=0.5, bullet_color=self.T["accent"], min_size=12)

    def chart(self, slide, spec, takeaways):
        kinds = {"column": XL_CHART_TYPE.COLUMN_CLUSTERED, "bar": XL_CHART_TYPE.BAR_CLUSTERED,
                 "line": XL_CHART_TYPE.LINE_MARKERS, "pie": XL_CHART_TYPE.PIE,
                 "doughnut": XL_CHART_TYPE.DOUGHNUT, "stacked": XL_CHART_TYPE.COLUMN_STACKED,
                 "area": XL_CHART_TYPE.AREA}
        kind = str(spec.get("type", "column")).lower()
        cats = [str(c) for c in spec.get("categories", [])][:12]
        series = spec.get("series") or []
        data = CategoryChartData()
        data.categories = cats
        for s in series[:6]:
            vals = []
            for v in list(s.get("values", []))[: len(cats)]:
                try:
                    vals.append(float(v))
                except (TypeError, ValueError):
                    vals.append(0.0)
            vals += [0.0] * (len(cats) - len(vals))
            data.add_series(str(s.get("name", "Series")), vals)
        circular = kind in ("pie", "doughnut")
        cw = 7.7 if takeaways else AW
        gf = slide.shapes.add_chart(kinds.get(kind, kinds["column"]), Inches(AX), Inches(AY), Inches(cw),
                                    Inches(AH), data)
        ch = gf.chart
        ch.font.size, ch.font.name = Pt(12), BODY_FONT
        ch.font.color.rgb = rgb(self.T["text"])
        if spec.get("title"):
            ch.has_title = True
            ch.chart_title.text_frame.text = str(spec["title"])
            r = ch.chart_title.text_frame.paragraphs[0].runs[0]
            r.font.size, r.font.bold, r.font.name = Pt(14), True, BODY_FONT
            r.font.color.rgb = rgb(self.T["dark"])
        else:
            ch.has_title = False
        ch.has_legend = circular or len(series) > 1
        if ch.has_legend:
            ch.legend.position = XL_LEGEND_POSITION.BOTTOM
            ch.legend.include_in_layout = False
            ch.legend.font.size = Pt(12)
        palette = self.T["chart"]
        plot = ch.plots[0]
        if circular:
            light = [palette[k] for k in (0, 2, 3, 5, 4, 1)]
            for i in range(len(cats)):
                pt = plot.series[0].points[i]
                pt.format.fill.solid()
                pt.format.fill.fore_color.rgb = rgb(light[i % len(light)])
        else:
            for i, s in enumerate(plot.series):
                col = rgb(palette[i % len(palette)])
                if kind == "line":
                    s.format.line.color.rgb, s.format.line.width, s.smooth = col, Pt(3), False
                    s.marker.style, s.marker.size = XL_MARKER_STYLE.CIRCLE, 8
                    s.marker.format.fill.solid()
                    s.marker.format.fill.fore_color.rgb = col
                    s.marker.format.line.color.rgb = col
                else:
                    s.format.fill.solid()
                    s.format.fill.fore_color.rgb = col
            if kind in ("column", "bar", "stacked"):
                plot.gap_width = 60
            va, ca = ch.value_axis, ch.category_axis
            va.major_gridlines.format.line.color.rgb = rgb("E5E7EB")
            va.format.line.fill.background()
            ca.format.line.color.rgb = rgb("CBD5E1")
            va.tick_labels.font.size = ca.tick_labels.font.size = Pt(12)
            va.tick_labels.font.color.rgb = ca.tick_labels.font.color.rgb = rgb(self.T["muted"])
        if len(cats) * max(1, len(series)) <= 18:
            plot.has_data_labels = True
            dl = plot.data_labels
            dl.font.size, dl.font.bold = Pt(13), True
            dl.font.color.rgb = rgb("1F2937")
            if circular:
                dl.show_percentage, dl.show_value, dl.number_format = True, False, "0%"
                dl.number_format_is_linked = False
            else:
                dl.show_value = True
                dl.position = (XL_LABEL_POSITION.CENTER if kind in ("stacked", "area")
                               else XL_LABEL_POSITION.ABOVE if kind == "line" else XL_LABEL_POSITION.OUTSIDE_END)
        if takeaways:
            px, pw = AX + cw + 0.3, AW - cw - 0.3
            add_box(slide, px, AY, pw, AH, fill=self.T["card"])
            self.text(slide, px + 0.3, AY + 0.3, pw - 0.6, 0.4, "Key takeaways", size=16, bold=True,
                      color=self.T["accent"])
            self.text(slide, px + 0.3, AY + 0.9, pw - 0.6, AH - 1.2, takeaways[:4], size=15, bullets=True,
                      gap=0.6, bullet_color=self.T["accent"], min_size=11)

    def table(self, slide, spec):
        headers = [clean(h) for h in spec.get("headers", [])][:6]
        rows = [[clean(c) for c in r][: len(headers)] for r in spec.get("rows", [])][:8]
        if not headers:
            return
        ncol, nrow = len(headers), len(rows) + 1
        row_h = min(0.8, AH / nrow)
        gf = slide.shapes.add_table(nrow, ncol, Inches(AX), Inches(AY), Inches(AW), Inches(row_h * nrow))
        tbl = gf.table
        tbl.first_row, tbl.horz_banding = True, False
        weights = [min(40, max(8, len(headers[c]), *(len(r[c]) for r in rows if c < len(r)))) for c in range(ncol)]
        for c, wgt in enumerate(weights):
            tbl.columns[c].width = Inches(AW * wgt / sum(weights))
        size = 16 if nrow <= 6 else 13
        for r in range(nrow):
            tbl.rows[r].height = Inches(row_h)
            for c in range(ncol):
                cell = tbl.cell(r, c)
                val = headers[c] if r == 0 else (rows[r - 1][c] if c < len(rows[r - 1]) else "")
                cell.fill.solid()
                cell.fill.fore_color.rgb = rgb(self.T["dark"] if r == 0 else (self.T["card"] if r % 2 == 0 else "FFFFFF"))
                cell.vertical_anchor = MSO_ANCHOR.MIDDLE
                cell.margin_left = cell.margin_right = Inches(0.15)
                tf = cell.text_frame
                tf.word_wrap = True
                tf.text = truncate_to(val, 70)
                f = tf.paragraphs[0].runs[0].font
                f.name, f.size, f.bold = BODY_FONT, Pt(size), r == 0
                f.color.rgb = rgb("FFFFFF" if r == 0 else self.T["text"])

    def agenda(self, titles):
        slide = self.content_slide("Agenda")
        items = [(t, "") for t in titles[:10]]
        self.rows(slide, items, AX, AY, AW, AH, cols=2 if len(items) > 5 else 1, max_row_h=0.85)
        return slide

    # ---------- layout selection + dispatch ----------

    def pick_layout(self, s, items):
        explicit = str(s.get("layout") or "auto").lower()
        alias = {"bullets": "rows", "list": "rows", "timeline": "process", "comparison": "compare",
                 "kpi": "stats", "quote": "spotlight", "divider": "section"}
        explicit = alias.get(explicit, explicit)
        if s.get("chart"):
            return "chart"
        if s.get("table"):
            return "table"
        if s.get("stats"):
            return "stats"
        if s.get("steps"):
            return "process"
        if s.get("left") and s.get("right"):
            return "compare"
        if explicit in ("cards", "rows", "split", "stats", "process", "section", "spotlight"):
            return explicit
        if not items:
            return "section"
        parsed = [parse_stat(i) for i in items]
        if 2 <= len(items) <= 4 and all(parsed):
            return "stats"
        if len(items) == 1:
            return "spotlight"
        choice = ("cards", "rows", "split")[self.variety % 3]
        self.variety += 1
        if choice == "split" and len(items) < 3:
            choice = "cards"
        return choice

    def add_content(self, s):
        title = clean(s.get("title")) or "Untitled"
        items = normalize_items(s.get("content"))
        layout = self.pick_layout(s, items)
        pairs = [split_item(i) for i in items]

        if layout == "section":
            slide = self.section_slide(title, clean(s.get("subtitle")) or (items[0] if items else ""))
            return self.notes(slide, s.get("notes"))

        slide = self.content_slide(title)
        if layout == "chart":
            self.chart(slide, s["chart"], items)
        elif layout == "table":
            self.table(slide, s["table"])
        elif layout == "stats":
            raw = s.get("stats") or []
            stats = []
            for r in raw[:4]:
                if isinstance(r, dict):
                    stats.append((clean(r.get("value")), clean(r.get("label") or r.get("description"))))
                elif parse_stat(clean(r)):
                    stats.append(parse_stat(clean(r)))
            if not stats:
                stats = [p for p in (parse_stat(i) for i in items) if p][:4]
                items = []
            if stats:
                self.stats(slide, stats, items if s.get("stats") else None)
            else:
                self.cards(slide, pairs[:6])
        elif layout == "process":
            steps = []
            for r in s.get("steps") or []:
                if isinstance(r, dict):
                    steps.append((clean(r.get("title") or r.get("heading")), clean(r.get("description") or r.get("text"))))
                else:
                    steps.append(split_item(clean(r)))
            self.process(slide, (steps or pairs)[:8])
        elif layout == "compare":
            def side(d):
                pts = normalize_items(d.get("points") or d.get("content"), 6)
                return {"heading": clean(d.get("heading") or d.get("title") or ""), "points": pts}
            self.compare(slide, side(s["left"]), side(s["right"]))
        elif layout == "spotlight":
            self.spotlight(slide, items[0] if items else title)
        elif layout == "split" and len(items) >= 3:
            self.split(slide, items)
        elif layout == "rows" or layout == "split":
            self.rows(slide, pairs[:6], AX, AY, AW, AH, cols=2 if len(pairs) > 4 else 1)
        else:
            self.cards(slide, pairs[:6])
        self.notes(slide, s.get("notes"))

    def save(self, path, author=""):
        cp = self.prs.core_properties
        cp.title, cp.author = self.title, author or ""
        self.prs.save(path)


# =========================
# CLEAN-UP OF OLD FILES
# =========================

def purge_old_files():
    cutoff = time.time() - RETENTION_HOURS * 3600
    for name in os.listdir(OUTPUT_FOLDER):
        p = os.path.join(OUTPUT_FOLDER, name)
        try:
            if name.endswith(".pptx") and os.path.getmtime(p) < cutoff:
                os.remove(p)
        except OSError:
            pass


# =========================
# ROUTES
# =========================

@app.route("/")
def home():
    return "PPT Generator Running"


@app.route("/health")
def health():
    return jsonify({"status": "ok"})


@app.route("/createppt", methods=["POST"])
def create_ppt():
    data = request.get_json(silent=True)
    if not isinstance(data, dict):
        return jsonify({"status": "error", "message": "Request body must be JSON"}), 400

    slides = data.get("slides", [])
    if isinstance(slides, str):  # Copilot Studio sometimes sends the array as a string
        try:
            slides = json.loads(slides)
        except ValueError:
            slides = []
    slides = [s for s in slides if isinstance(s, dict)][:MAX_CONTENT_SLIDES]
    if not slides:
        return jsonify({"status": "error", "message": "'slides' must be a non-empty list"}), 400

    title = clean(data.get("presentationTitle")) or "Presentation"
    footer = data.get("footer")

    try:
        deck = Deck(title, data.get("theme", "midnight"), None if footer is None else clean(footer))
        meta = " · ".join(x for x in (clean(data.get("presenter")), clean(data.get("date"))) if x)
        deck.title_slide(clean(data.get("subtitle")), clean(data.get("organization")), meta)

        content_titles = [clean(s.get("title")) for s in slides
                          if str(s.get("layout", "")).lower() not in ("section", "divider") and clean(s.get("title"))]
        if data.get("agenda", True) and len(content_titles) >= 4:
            deck.agenda(content_titles)

        for s in slides:
            deck.add_content(s)

        if data.get("closing", True):
            deck.closing_slide(clean(data.get("closingTitle")) or "Thank You",
                               clean(data.get("closingSubtitle")) or "Questions & discussion")

        purge_old_files()
        filename = f"{uuid.uuid4()}.pptx"
        deck.save(os.path.join(OUTPUT_FOLDER, filename), clean(data.get("presenter")))
    except Exception as exc:  # never leak a stack trace to the agent, but tell it what failed
        app.logger.exception("PPT generation failed")
        return jsonify({"status": "error", "message": f"Could not generate deck: {exc}"}), 500

    base = (os.environ.get("PUBLIC_BASE_URL") or request.url_root).rstrip("/") + "/"
    return jsonify({
        "status": "success",
        "fileName": filename,
        "downloadUrl": base + "download/" + filename,
        "slideCount": len(deck.prs.slides),
    })


@app.route("/download/<filename>")
def download(filename):
    return send_from_directory(OUTPUT_FOLDER, filename, as_attachment=True)


# =========================
# MAIN
# =========================

if __name__ == "__main__":
    app.run(host="0.0.0.0", port=int(os.environ.get("PORT", 5000)))
